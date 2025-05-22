from flask import Flask, request, jsonify, send_file
from werkzeug.utils import secure_filename
import os
import uuid
import cv2
from PIL import Image
from pathlib import Path
from fpdf import FPDF
from docx import Document
import json
from flask_cors import CORS

from document_pipeline import DocumentRectifierOCR  # Assuming it's saved as document_pipeline.py
from ocr_utils import perform_ocr  # Assuming perform_ocr is defined elsewhere
from rectify import rectify_single_image  # Assuming rectify_single_image is defined in rectify.py

# Flask app initialization
app = Flask(__name__)

CORS(app)  # Enable CORS for all routes
# Set base directories
BASE_DIR = Path(__file__).resolve().parent.parent  # Go up from /backend/ to project root
STORAGE_DIR = BASE_DIR / "storage"
UPLOAD_FOLDER = STORAGE_DIR / "uploads"
RESULT_FOLDER = STORAGE_DIR / "results"

# Ensure folders exist
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESULT_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = str(UPLOAD_FOLDER)

sessions = {}  # session_id -> dictionary with 'images' and 'ocr' results

# Initialize document rectifier and OCR processor
document_processor = DocumentRectifierOCR()

@app.route('/uploads/<filename>')
def serve_processed_image(filename):
    """Serve processed image from uploads folder."""
    file_path = UPLOAD_FOLDER / filename
    if not file_path.exists():
        return jsonify({"error": "File not found"}), 404
    return send_file(file_path)

@app.route('/api/ocr', methods=['POST'])
def ocr():
    """Perform OCR on all images in a session, individually."""
    print("OCR endpoint called")
    
    if request.is_json:
        data = request.get_json()
        session_id = data.get("sessionId") or data.get("session_id")
    else:
        session_id = request.form.get("sessionId") or request.form.get("session_id")
    
    if not session_id or session_id not in sessions:
        return jsonify({"error": "Invalid session ID"}), 400

    # Make sure ocr key exists
    sessions[session_id]["ocr"] = {}

    response_data = []
    for img in sessions[session_id]["images"]:
        img_id = img['id']
        img_path = img['path']
        image = cv2.imread(str(img_path))

        text, conf = document_processor.extract_text(image)
        char_count = len(text)
        avg_conf = conf if char_count else 0

        # Save text result per image
        output_path = RESULT_FOLDER / f"{img_id}.txt"
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)

        sessions[session_id]["ocr"][img_id] = {
            "text": text,
            "confidence": round(avg_conf, 2),
            "characters": char_count,
            "path": output_path
        }

        response_data.append({
            "imageId": img_id,
            "textPreview": text[:1000],
            "characterCount": char_count,
            "confidenceScore": round(avg_conf, 2),
            "originalName": img.get("name", "Image")
        })

    return jsonify(response_data)

@app.route('/api/download', methods=['GET'])
def download():
    """Download OCR result for a single image."""
    image_id = request.args.get("id")
    file_format = request.args.get("format", "txt").lower()

    # Search through sessions
    found = None
    for session in sessions.values():
        if "ocr" in session and image_id in session["ocr"]:
            found = session["ocr"][image_id]
            break

    if not found:
        return jsonify({"error": "OCR result not found"}), 404

    text = found["text"]
    output_path = RESULT_FOLDER / f"{image_id}.{file_format}"

    if file_format == "txt":
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)
    elif file_format == "json":
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump({"text": text}, f, ensure_ascii=False)
    elif file_format == "pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)
        for line in text.splitlines():
            pdf.cell(200, 10, txt=line, ln=True)
        pdf.output(str(output_path))
    elif file_format == "docx":
        doc = Document()
        doc.add_paragraph(text)
        doc.save(str(output_path))
    else:
        return jsonify({"error": "Unsupported format"}), 400

    return send_file(output_path, as_attachment=True)


@app.route('/api/preprocess', methods=['POST'])
def preprocess_images():
    """Upload images, rectify them, and save."""
    print("=== /api/preprocess endpoint called ===")
    
    if 'images' not in request.files:
        print("Error: No 'images' in request.files")
        print(f"request.files keys: {list(request.files.keys())}")
        print(f"request.form: {request.form}")
        return jsonify({'error': 'No images uploaded'}), 400

    session_id = str(uuid.uuid4())
    print(f"Created new session: {session_id}")
    
    sessions[session_id] = {"images": [], "ocr": {}}
    uploaded_files = request.files.getlist('images')
    print(f"Received {len(uploaded_files)} files")

    for i, file in enumerate(uploaded_files):
        if file.filename == '':
            print(f"File {i} has empty filename, skipping")
            continue

        print(f"Processing file {i}: {file.filename}")
        temp_path = UPLOAD_FOLDER / f"temp_{uuid.uuid4().hex}.png"
        file.save(str(temp_path))
        print(f"Saved original image: {temp_path}")

        try:
            print(f"Attempting to rectify image {file.filename}")
            scanned = rectify_single_image(str(temp_path))
            if scanned is None:
                print(f"Rectification returned None for {file.filename}")
                raise ValueError("Rectification returned None.")

            print(f"Successfully processed image for {file.filename}")

            processed_id = uuid.uuid4().hex
            processed_path = UPLOAD_FOLDER / f"{processed_id}.png"
            scanned.save(str(processed_path))
            print(f"Processed image saved at: {processed_path}")

            # Store image information in the session dictionary
            sessions[session_id]["images"].append({
                "id": processed_id,
                "path": processed_path,
                "name": file.filename
            })

            os.remove(temp_path)
            print(f"Temporary file removed: {temp_path}")

        except Exception as e:
            print(f"Error processing image {file.filename}: {e}")
            import traceback
            traceback.print_exc()
            return jsonify({"error": f"Failed to preprocess {file.filename}: {str(e)}"}), 500

    if not sessions[session_id]["images"]:
        print("No images were processed successfully")
        return jsonify({"error": "No images processed successfully"}), 400

    processed_images = [
        {
            "id": img["id"],
            "url": f"/uploads/{img['id']}.png",
            "name": img["name"]
        }
        for img in sessions[session_id]["images"]
    ]

    response_data = {
        "session_id": session_id,
        "images": processed_images
    }
    
    print(f"Returning processed images for session {session_id}")
    print(f"Response payload: {json.dumps(response_data)}")
    
    return jsonify(response_data)


if __name__ == "__main__":
    app.run(debug=True)
